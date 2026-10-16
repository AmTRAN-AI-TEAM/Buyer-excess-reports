from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from math import isfinite
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
DEFAULT_DEMAND_LABEL = "WO外demand"
PART_NO_HEADERS = ("Part No", "PartNo", "Part Number", "COMPONENT_ITEM")
ITEM_PART_HEADERS = ("Item Number", "Part No", "Part No.", "COMPONENT_ITEM")


@dataclass(frozen=True)
class ShortageSource:
    path: Path
    sheet_name: str


@dataclass(frozen=True)
class DemandBucket:
    label: str
    columns: tuple[int, ...]


@dataclass(frozen=True)
class ShortageColumns:
    header_row: int
    part_no: int
    description: int
    planner: int
    buyer: int
    lt: int
    moq: int
    demand_buckets: tuple[DemandBucket, ...]
    over_shortage: int | None
    hld: int | None
    overshortage_with_hld: int | None
    po_remain: int | None


@dataclass(frozen=True)
class ShortageRecord:
    part_no: str
    description: Any
    planner: Any
    buyer: Any
    lt: Any
    moq: Any
    overshortage: float
    wo_demands: tuple[float, ...]
    demand_labels: tuple[str, ...]
    open_po: float
    source_row: int

    @property
    def wo_demand(self) -> float:
        return sum(self.wo_demands)

    def wo_demand_at(self, index: int) -> float:
        if index >= len(self.wo_demands):
            return 0.0
        return self.wo_demands[index]


@dataclass(frozen=True)
class ItemMetadata:
    customer: str | None
    model: str | None


@dataclass(frozen=True)
class PriceCandidate:
    price: float
    allocation: float
    start_date: datetime
    row_number: int


@dataclass(frozen=True)
class ItemWorkbookData:
    metadata_by_part: dict[str, ItemMetadata]
    prices_by_part: dict[str, float]

    def customer_summary(self, parts: list[str]) -> str | None:
        return _join_limited(
            [
                value
                for part in parts
                for value in [_metadata_value(self.metadata_by_part, part, "customer")]
                if value
            ]
        )

    def model_summary(self, parts: list[str]) -> str | None:
        return _join_limited(
            [
                value
                for part in parts
                for value in [_metadata_value(self.metadata_by_part, part, "model")]
                if value
            ]
        )

    def price_for(self, part_no: str) -> float | None:
        return self.prices_by_part.get(part_no)


def excel_files(input_dir: Path) -> list[Path]:
    return sorted(
        p
        for p in input_dir.iterdir()
        if p.is_file()
        and p.suffix.lower() in EXCEL_EXTENSIONS
        and not p.name.startswith("~$")
    )


def find_bom_file(input_dir: Path) -> Path:
    candidates = [p for p in excel_files(input_dir) if "bom" in p.name.lower()]
    if len(candidates) == 1:
        return candidates[0]

    if not candidates:
        raise FileNotFoundError(
            f"Cannot find an Excel file with BOM in its name in {input_dir}"
        )

    names = ", ".join(p.name for p in candidates)
    raise ValueError(f"Found multiple BOM Excel files: {names}")


def find_item_file(input_dir: Path) -> Path:
    candidates = [p for p in excel_files(input_dir) if "item" in p.name.lower()]
    if len(candidates) == 1:
        return candidates[0]

    if not candidates:
        raise FileNotFoundError(
            f"Cannot find an Excel file with Item in its name in {input_dir}"
        )

    names = ", ".join(p.name for p in candidates)
    raise ValueError(f"Found multiple Item Excel files: {names}")


def find_price_file(input_dir: Path) -> Path:
    candidates = [p for p in excel_files(input_dir) if "price" in p.name.lower()]
    if len(candidates) == 1:
        return candidates[0]

    if not candidates:
        raise FileNotFoundError(
            f"Cannot find an Excel file with Price in its name in {input_dir}"
        )

    names = ", ".join(p.name for p in candidates)
    raise ValueError(f"Found multiple Price Excel files: {names}")


def find_shortage_source(input_dir: Path, bom_file: Path) -> ShortageSource:
    matches: list[ShortageSource] = []
    for workbook_path in excel_files(input_dir):
        if workbook_path.resolve() == bom_file.resolve():
            continue

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        for sheet_name in workbook.sheetnames:
            if "shortage" in sheet_name.lower():
                matches.append(ShortageSource(workbook_path, sheet_name))
        workbook.close()

    if len(matches) == 1:
        return matches[0]

    if not matches:
        raise FileNotFoundError(
            f"Cannot find a worksheet with 'shortage' in its name under {input_dir}"
        )

    details = ", ".join(f"{m.path.name}:{m.sheet_name}" for m in matches)
    raise ValueError(f"Found multiple shortage worksheets; please keep only one: {details}")


def load_shortage_records(source: ShortageSource) -> list[ShortageRecord]:
    workbook = load_workbook(source.path, read_only=True, data_only=True)
    worksheet = workbook[source.sheet_name]
    columns = detect_shortage_columns(worksheet)
    demand_labels = tuple(bucket.label for bucket in columns.demand_buckets)

    records: list[ShortageRecord] = []
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=columns.header_row + 1, values_only=True),
        start=columns.header_row + 1,
    ):
        part_no = _cell(row, columns.part_no)
        if part_no in (None, ""):
            continue

        overshortage = _overshortage_value(row, columns)
        wo_demands = _wo_demand_values(row, columns)
        open_po = _number(_cell(row, columns.po_remain)) if columns.po_remain else 0.0

        records.append(
            ShortageRecord(
                part_no=str(part_no).strip(),
                description=_cell(row, columns.description),
                planner=_cell(row, columns.planner),
                buyer=_cell(row, columns.buyer),
                lt=_cell(row, columns.lt),
                moq=_cell(row, columns.moq),
                overshortage=overshortage,
                wo_demands=wo_demands,
                demand_labels=demand_labels,
                open_po=open_po,
                source_row=row_index,
            )
        )

    workbook.close()
    return records


def detect_shortage_columns(worksheet: Worksheet) -> ShortageColumns:
    header_row = _find_header_row(worksheet)
    headers = {
        _normalize_header(worksheet.cell(header_row, column).value): column
        for column in range(1, worksheet.max_column + 1)
    }

    def require(*names: str) -> int:
        for name in names:
            normalized = _normalize_header(name)
            if normalized in headers:
                return headers[normalized]
        raise ValueError(
            f"Cannot find required shortage column {names} on row {header_row}"
        )

    def optional(*names: str) -> int | None:
        for name in names:
            normalized = _normalize_header(name)
            if normalized in headers:
                return headers[normalized]
        return None

    part_no = require(*PART_NO_HEADERS)
    leading_demand_columns = tuple(
        column
        for column in range(1, part_no)
        if _is_demand_header(worksheet.cell(header_row, column).value)
    )
    mps_columns = tuple(
        column
        for column in range(1, worksheet.max_column + 1)
        if _normalize_header(worksheet.cell(header_row - 1, column).value) == "mps"
    )
    demand_buckets = _demand_buckets(
        worksheet,
        header_row,
        leading_demand_columns,
        mps_columns,
    )

    return ShortageColumns(
        header_row=header_row,
        part_no=part_no,
        description=require("DESCRIPTION", "Description", "Desc"),
        planner=require("PLANNER", "Planner"),
        buyer=require("Buyer", "GSD"),
        lt=require("LT"),
        moq=require("MOQ"),
        demand_buckets=demand_buckets,
        over_shortage=optional("OVER_SHORTAGE", "Overshortage"),
        hld=optional("HLD"),
        overshortage_with_hld=optional("Overshortage1"),
        po_remain=optional("PO_REMAIN", "PO REMAIN", "Open PO"),
    )


def load_item_workbook_data(
    item_path: Path,
    *,
    price_path: Path | None = None,
    price_headers: tuple[str, ...] = ("Uprice",),
    preferred_cost_sheet: str | None = None,
) -> ItemWorkbookData:
    workbook = load_workbook(item_path, read_only=True, data_only=True)
    try:
        item_sheet = _find_item_sheet(workbook)
        metadata = _load_item_metadata(workbook[item_sheet])
        preferred_prices: dict[str, float] = {}
        cost_sheet = _find_sheet_or_none(workbook.sheetnames, preferred_cost_sheet) if preferred_cost_sheet else None
        if cost_sheet is not None:
            preferred_prices = {
                part: price
                for part, price in _load_prices(
                    workbook[cost_sheet], ("US Unit Cost",), part_headers=("Item",),
                ).items()
                if isfinite(price)
            }
    finally:
        workbook.close()

    price_workbook_path = price_path or item_path
    price_workbook = load_workbook(price_workbook_path, read_only=True, data_only=True)
    try:
        try:
            price_sheet = _find_price_sheet(price_workbook, price_headers)
        except ValueError:
            if not preferred_prices:
                raise
            prices = {}
        else:
            prices = _load_prices(price_workbook[price_sheet], price_headers)
    finally:
        price_workbook.close()

    # A numeric zero is a valid cost and must override the fallback price too.
    prices.update(preferred_prices)
    return ItemWorkbookData(metadata_by_part=metadata, prices_by_part=prices)


def _find_header_row(worksheet: Worksheet) -> int:
    for row in range(1, min(worksheet.max_row, 30) + 1):
        values = [
            _normalize_header(worksheet.cell(row, column).value)
            for column in range(1, min(worksheet.max_column, 40) + 1)
        ]
        if any(_normalize_header(name) in values for name in PART_NO_HEADERS) and (
            "description" in values or "desc" in values
        ):
            return row
    raise ValueError("Cannot identify the shortage header row")


def _overshortage_value(row: tuple[Any, ...], columns: ShortageColumns) -> float:
    if columns.overshortage_with_hld:
        value = _cell(row, columns.overshortage_with_hld)
        if value is not None:
            return _number(value)

    over_shortage = _number(_cell(row, columns.over_shortage)) if columns.over_shortage else 0
    hld = _number(_cell(row, columns.hld)) if columns.hld else 0
    return over_shortage + hld


def _wo_demand_values(row: tuple[Any, ...], columns: ShortageColumns) -> tuple[float, ...]:
    return tuple(
        sum(_number(_cell(row, column)) for column in bucket.columns)
        for bucket in columns.demand_buckets
    )


def _demand_buckets(
    worksheet: Worksheet,
    header_row: int,
    leading_demand_columns: tuple[int, ...],
    mps_columns: tuple[int, ...],
) -> tuple[DemandBucket, ...]:
    if leading_demand_columns:
        labels = _unique_labels(
            [
                _demand_label(
                    _demand_label_source(worksheet, header_row, column),
                    index,
                )
                for index, column in enumerate(leading_demand_columns, start=1)
            ]
        )
        return tuple(
            DemandBucket(label=label, columns=(column,))
            for label, column in zip(labels, leading_demand_columns)
        )

    if mps_columns:
        return (DemandBucket(label=DEFAULT_DEMAND_LABEL, columns=mps_columns),)

    return (DemandBucket(label=DEFAULT_DEMAND_LABEL, columns=()),)


def _demand_label(value: Any, index: int) -> str:
    text = _text_or_none(value)
    if text:
        return text
    if index == 1:
        return DEFAULT_DEMAND_LABEL
    return f"{DEFAULT_DEMAND_LABEL} {index}"


def _demand_label_source(
    worksheet: Worksheet,
    header_row: int,
    column: int,
) -> Any:
    if header_row > 1:
        previous_value = _text_or_none(worksheet.cell(header_row - 1, column).value)
        if previous_value:
            return previous_value
    return worksheet.cell(header_row, column).value


def _is_demand_header(value: Any) -> bool:
    normalized = _normalize_header(value)
    if normalized == _normalize_header(DEFAULT_DEMAND_LABEL):
        return True

    text = str(value or "").casefold()
    return "工单外" in text and "demand" in text


def _unique_labels(labels: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    unique: list[str] = []
    for label in labels:
        count = seen.get(label, 0) + 1
        seen[label] = count
        unique.append(label if count == 1 else f"{label} {count}")
    return unique


def _load_item_metadata(worksheet: Worksheet) -> dict[str, ItemMetadata]:
    header_row_number = _find_table_header_row(
        worksheet,
        (ITEM_PART_HEADERS, ("Mps Customer",), ("Mps Model",)),
    )
    if header_row_number is None:
        return {}

    rows = worksheet.iter_rows(min_row=header_row_number, values_only=True)
    header_row = next(rows)
    headers = _header_map(header_row)
    item_number = _require_header(headers, *ITEM_PART_HEADERS)
    customer = _require_header(headers, "Mps Customer")
    model = _require_header(headers, "Mps Model")

    metadata: dict[str, ItemMetadata] = {}
    for row in rows:
        part_no = _text_or_none(_cell_by_index(row, item_number))
        if not part_no:
            continue
        metadata[part_no] = ItemMetadata(
            customer=_text_or_none(_cell_by_index(row, customer)),
            model=_text_or_none(_cell_by_index(row, model)),
        )
    return metadata


def _load_prices(
    worksheet: Worksheet,
    price_headers: tuple[str, ...],
    *,
    part_headers: tuple[str, ...] = PART_NO_HEADERS,
) -> dict[str, float]:
    header_row_number = _find_table_header_row(
        worksheet,
        (part_headers, price_headers),
    )
    if header_row_number is None:
        return {}

    rows = worksheet.iter_rows(min_row=header_row_number, values_only=True)
    header_row = next(rows)
    headers = _header_map(header_row)
    part_no_column = _require_header(headers, *part_headers)
    uprice_column = _require_header(headers, *price_headers)
    allocation_column = _optional_header(headers, "Allocation")
    start_date_column = _optional_header(headers, "Start Date")

    candidates: dict[str, PriceCandidate] = {}
    for row_number, row in enumerate(rows, start=header_row_number + 1):
        part_no = _text_or_none(_cell_by_index(row, part_no_column))
        price = _number_or_none(_cell_by_index(row, uprice_column))
        if not part_no or price is None:
            continue

        candidate = PriceCandidate(
            price=price,
            allocation=_number(_cell_by_index(row, allocation_column)),
            start_date=_date_value(_cell_by_index(row, start_date_column)),
            row_number=row_number,
        )
        current = candidates.get(part_no)
        if current is None or _price_rank(candidate) > _price_rank(current):
            candidates[part_no] = candidate

    return {part_no: candidate.price for part_no, candidate in candidates.items()}


def _find_item_sheet(workbook: Any) -> str:
    exact_sheet = _find_sheet_or_none(workbook.sheetnames, "Item")
    if exact_sheet is not None:
        return exact_sheet
    return _find_sheet_with_headers(
        workbook,
        (ITEM_PART_HEADERS, ("Mps Customer",), ("Mps Model",)),
        "Item metadata",
    )


def _find_price_sheet(workbook: Any, price_headers: tuple[str, ...]) -> str:
    exact_sheet = _find_sheet_or_none(workbook.sheetnames, "Price")
    if exact_sheet is not None:
        return exact_sheet
    return _find_sheet_with_headers(
        workbook,
        (PART_NO_HEADERS, price_headers),
        "Price",
    )


def _find_sheet_with_headers(
    workbook: Any,
    required_header_groups: tuple[tuple[str, ...], ...],
    description: str,
) -> str:
    for sheet_name in workbook.sheetnames:
        if _find_table_header_row(workbook[sheet_name], required_header_groups) is not None:
            return sheet_name
    raise ValueError(f"Cannot find required {description} sheet")


def _find_table_header_row(
    worksheet: Worksheet,
    required_header_groups: tuple[tuple[str, ...], ...],
    *,
    max_rows: int = 30,
) -> int | None:
    for row_number, row in enumerate(
        worksheet.iter_rows(max_row=min(worksheet.max_row, max_rows), values_only=True),
        start=1,
    ):
        headers = _header_map(row)
        if all(
            any(_normalize_header(header) in headers for header in group)
            for group in required_header_groups
        ):
            return row_number
    return None


def _price_rank(candidate: PriceCandidate) -> tuple[float, datetime, int]:
    return (candidate.allocation, candidate.start_date, -candidate.row_number)


def _header_map(header_row: tuple[Any, ...]) -> dict[str, int]:
    return {
        _normalize_header(value): index
        for index, value in enumerate(header_row)
        if value not in (None, "")
    }


def _require_header(headers: dict[str, int], *names: str) -> int:
    for name in names:
        normalized = _normalize_header(name)
        if normalized in headers:
            return headers[normalized]
    raise ValueError(f"Cannot find required column {names}")


def _optional_header(headers: dict[str, int], *names: str) -> int | None:
    for name in names:
        normalized = _normalize_header(name)
        if normalized in headers:
            return headers[normalized]
    return None


def _cell(row: tuple[Any, ...], one_based_column: int | None) -> Any:
    if one_based_column is None:
        return None
    index = one_based_column - 1
    if index >= len(row):
        return None
    return row[index]


def _cell_by_index(row: tuple[Any, ...], zero_based_column: int | None) -> Any:
    if zero_based_column is None or zero_based_column >= len(row):
        return None
    return row[zero_based_column]


def _number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def _number_or_none(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _date_value(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    return datetime.min


def _find_sheet(sheet_names: list[str], expected_name: str) -> str:
    sheet_name = _find_sheet_or_none(sheet_names, expected_name)
    if sheet_name is not None:
        return sheet_name
    raise ValueError(f"Cannot find required sheet {expected_name}")


def _find_sheet_or_none(sheet_names: list[str], expected_name: str) -> str | None:
    for sheet_name in sheet_names:
        if sheet_name.lower() == expected_name.lower():
            return sheet_name
    return None


def _metadata_value(
    metadata_by_part: dict[str, ItemMetadata],
    part_no: str,
    field_name: str,
) -> str | None:
    metadata = metadata_by_part.get(part_no)
    if metadata is None:
        return None
    return getattr(metadata, field_name)


def _join_limited(values: list[str], limit: int = 6) -> str | None:
    unique: list[str] = []
    for value in values:
        if value not in unique:
            unique.append(value)

    if not unique:
        return None
    if len(unique) <= limit:
        return "/".join(unique)
    return "/".join(unique[:limit]) + f"/+{len(unique) - limit} more"


def _text_or_none(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    return text or None


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return "".join(ch for ch in str(value).strip().lower() if ch.isalnum())
