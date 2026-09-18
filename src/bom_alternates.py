from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


@dataclass(frozen=True)
class BomAlternateSegment:
    row_number: int
    parts: tuple[str, ...]
    base_part: str
    customer: str | None
    model: str | None


@dataclass
class BomAlternates:
    segments_by_part: dict[str, tuple[BomAlternateSegment, ...]]
    first_seen_order: dict[str, int]
    customers_by_part: dict[str, tuple[str, ...]] = field(default_factory=dict)
    models_by_part: dict[str, tuple[str, ...]] = field(default_factory=dict)

    def ordered_group_for(
        self,
        part_no: str,
        known_parts: set[str] | None = None,
        active_parts: set[str] | None = None,
    ) -> tuple[str, ...]:
        segment = self._best_segment_for(part_no, known_parts or set(), active_parts or set())
        if segment is None:
            return (part_no,)
        return segment.parts

    def customer_summary(self, parts: list[str]) -> str | None:
        return _join_limited(_unique_values(self.customers_by_part, parts))

    def model_summary(self, parts: list[str]) -> str | None:
        return _join_limited(_unique_values(self.models_by_part, parts))

    def _best_segment_for(
        self,
        part_no: str,
        known_parts: set[str],
        active_parts: set[str],
    ) -> BomAlternateSegment | None:
        candidates = self.segments_by_part.get(part_no, ())
        if not candidates:
            return None

        return max(
            candidates,
            key=lambda segment: (
                segment.base_part == part_no,
                sum(part in active_parts for part in segment.parts),
                sum(part in known_parts for part in segment.parts),
                len(segment.parts),
                -segment.row_number,
            ),
        )


def load_bom_alternates(bom_path: Path) -> BomAlternates:
    workbook = load_workbook(bom_path, read_only=True, data_only=True)
    bom_sheet = _find_sheet(workbook.sheetnames, "BOM")
    bom_ws = workbook[bom_sheet]

    finished_goods = _load_finished_goods(workbook)
    first_seen: dict[str, int] = {}
    customers: dict[str, list[str]] = defaultdict(list)
    models: dict[str, list[str]] = defaultdict(list)
    segments: list[BomAlternateSegment] = []

    current_context: tuple[Any, Any, Any] | None = None
    current_base: str | None = None
    current_base_row: int | None = None
    current_parts: list[str] = []
    current_customer: str | None = None
    current_model: str | None = None
    order = 0

    for row_number, row in enumerate(bom_ws.iter_rows(min_row=2, values_only=True), start=2):
        top_assembly = row[0] if len(row) > 0 else None
        bill_level = row[2] if len(row) > 2 else None
        assembly_item = row[3] if len(row) > 3 else None
        item_seq = row[6] if len(row) > 6 else None
        component = row[7] if len(row) > 7 else None

        if component in (None, ""):
            continue

        part_no = str(component).strip()
        if part_no not in first_seen:
            first_seen[part_no] = order
            order += 1

        _collect_finished_good_metadata(
            part_no,
            top_assembly,
            finished_goods,
            customers,
            models,
        )

        context = (top_assembly, bill_level, assembly_item)
        if _is_replacement_marker(item_seq):
            if current_base and current_context == context:
                current_parts.append(part_no)
            continue

        _append_segment(
            segments,
            current_base_row,
            current_base,
            current_parts,
            current_customer,
            current_model,
        )
        current_context = context
        current_base = part_no
        current_base_row = row_number
        current_parts = [part_no]
        current_customer, current_model = _finished_good_metadata(
            top_assembly,
            finished_goods,
        )

    _append_segment(
        segments,
        current_base_row,
        current_base,
        current_parts,
        current_customer,
        current_model,
    )

    segments_by_part: dict[str, list[BomAlternateSegment]] = defaultdict(list)
    for segment in segments:
        for part_no in segment.parts:
            segments_by_part[part_no].append(segment)

    workbook.close()
    return BomAlternates(
        segments_by_part={
            part_no: tuple(part_segments)
            for part_no, part_segments in segments_by_part.items()
        },
        first_seen_order=first_seen,
        customers_by_part={key: tuple(values) for key, values in customers.items()},
        models_by_part={key: tuple(values) for key, values in models.items()},
    )


def _find_sheet(sheet_names: list[str], expected_name: str) -> str:
    for sheet_name in sheet_names:
        if sheet_name.lower() == expected_name.lower():
            return sheet_name
    raise ValueError(f"Cannot find required sheet {expected_name}")


def _load_finished_goods(workbook: Any) -> dict[str, tuple[Any, Any]]:
    sheet_name = next(
        (name for name in workbook.sheetnames if name.lower() == "成品料号"),
        None,
    )
    if sheet_name is None:
        return {}

    worksheet = workbook[sheet_name]
    mapping: dict[str, tuple[Any, Any]] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        customer = row[0] if len(row) > 0 else None
        finished_good = row[1] if len(row) > 1 else None
        model = row[2] if len(row) > 2 else None
        if finished_good not in (None, ""):
            mapping[str(finished_good).strip()] = (customer, model)
    return mapping


def _append_segment(
    segments: list[BomAlternateSegment],
    row_number: int | None,
    base_part: str | None,
    parts: list[str],
    customer: str | None,
    model: str | None,
) -> None:
    if row_number is None or base_part is None or len(parts) <= 1:
        return

    unique_parts: list[str] = []
    for part_no in parts:
        if part_no not in unique_parts:
            unique_parts.append(part_no)

    segments.append(
        BomAlternateSegment(
            row_number=row_number,
            parts=tuple(unique_parts),
            base_part=base_part,
            customer=customer,
            model=model,
        )
    )


def _finished_good_metadata(
    top_assembly: Any,
    finished_goods: dict[str, tuple[Any, Any]],
) -> tuple[str | None, str | None]:
    if top_assembly in (None, ""):
        return None, None

    customer, model = finished_goods.get(str(top_assembly).strip(), (None, None))
    return _text_or_none(customer), _text_or_none(model)


def _collect_finished_good_metadata(
    part_no: str,
    top_assembly: Any,
    finished_goods: dict[str, tuple[Any, Any]],
    customers: dict[str, list[str]],
    models: dict[str, list[str]],
) -> None:
    if top_assembly in (None, ""):
        return

    customer, model = finished_goods.get(str(top_assembly).strip(), (None, None))
    _append_unique(customers[part_no], customer)
    _append_unique(models[part_no], model)


def _is_replacement_marker(value: Any) -> bool:
    if value is None:
        return False
    normalized = str(value).strip().replace("*", "").upper()
    return normalized == "R"


def _append_unique(values: list[str], value: Any) -> None:
    text = _text_or_none(value)
    if text and text not in values:
        values.append(text)


def _text_or_none(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    return text or None


def _unique_values(source: dict[str, tuple[str, ...]], parts: list[str]) -> list[str]:
    values: list[str] = []
    for part_no in parts:
        for value in source.get(part_no, ()):
            if value not in values:
                values.append(value)
    return values


def _join_limited(values: list[str], limit: int = 6) -> str | None:
    if not values:
        return None
    if len(values) <= limit:
        return "/".join(values)
    return "/".join(values[:limit]) + f"/+{len(values) - limit} more"
