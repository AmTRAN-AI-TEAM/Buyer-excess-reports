from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Any

from openpyxl.utils import column_index_from_string, get_column_letter

from bom_alternates import BomAlternates
from report_history import PreviousExcessInputs, previous_value_for
from workbook_io import ItemWorkbookData, ShortageRecord


INFO_HEADERS = [
    "Desc",
    "Planner",
    "GSD",
    "Customer\n客户",
    "MODEL",
    "MODELRemark\n机种",
    "LT",
    "MOQ",
    "Price（USD)",
]
AVTC_INFO_HEADERS = [
    "Desc",
    "Planner",
    "GSD",
    "LT",
    "MOQ",
    "Price（USD)",
]
EXCESS_METRIC_HEADERS = [
    "Excess\nPOQty",
    "Excess\nPOAmount",
    "Excess\nstockQTY",
    "Excess\nstockAmount",
    "Excess\nTotalAMT",
]
AVTC_EXCESS_METRIC_HEADERS = [
    "Excess\nstockQTY",
    "Excess\nstockAmount",
    "Excess\nPOQty",
    "Excess\nPOAmount",
    "Excess\nTotalAMT",
]
OTHER_DEMAND_HEADER = "others demand"
PO_OUTSIDE_STOCK_HEADER = "PO外備料 QTY"
PO_OUTSIDE_QTY_HEADER = "Excess PO外備料QTY"
PO_OUTSIDE_AMOUNT_HEADER = "Excess  PO外備料AMT"
PO_OUTSIDE_HEADERS = [
    PO_OUTSIDE_STOCK_HEADER,
    PO_OUTSIDE_QTY_HEADER,
    PO_OUTSIDE_AMOUNT_HEADER,
]
TRAILING_HEADERS = [
    "Vendor",
    "reason&action",
    "Category\n分类",
    "Improve",
    "Previous Excess Total AMT",
    "是否呆",
]
AVTC_TRAILING_HEADERS = [
    "Vendor",
    "reason&action",
    "Customer\n客户",
    "MODEL",
    "MODELRemark\n机种",
    "Category\n分类",
    "Improve",
    "Previous Excess Total AMT",
    "是否呆",
]
DEFAULT_DEMAND_LABELS = ("WO外demand",)


@dataclass(frozen=True)
class ExcessLayout:
    info_headers: tuple[str, ...]
    trailing_headers: tuple[str, ...]
    excess_metric_headers: tuple[str, ...]
    include_excess_input_columns: bool = False
    blank_columns_before_excess: int = 0
    after_first_excess_headers: tuple[str, ...] = ()


@dataclass(frozen=True)
class ExcessBuildResult:
    rows: list[list[Any]]
    row1_headers: list[Any]
    row2_headers: list[str]
    part_column_count: int
    input_part_count: int
    output_row_count: int
    demand_labels: tuple[str, ...]
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class PreviousExcessTotals:
    totals_by_part: dict[str, float]
    totals_by_group: dict[tuple[str, ...], float]

    def total_for(self, parts: list[str]) -> float | None:
        return previous_value_for(parts, self.totals_by_part, self.totals_by_group)


@dataclass(frozen=True)
class _PreparedExcessRow:
    parts: list[str]
    main_record: ShortageRecord
    price: float | None
    overshortage_values: list[int]
    wo_values_by_period: list[list[int]]
    open_po_values: list[int]
    overshortage_total: int
    wo_totals: list[int]
    open_po_total: int
    customer: str | None
    model: str | None
    model_remark: str | None
    carried_values: dict[str, Any] = field(default_factory=dict)


def build_excess_rows(
    shortage_records: list[ShortageRecord],
    alternates: BomAlternates,
    item_data: ItemWorkbookData | None = None,
    *,
    include_excess_input_columns: bool = False,
    previous_excess_totals: PreviousExcessTotals | None = None,
    previous_excess_inputs: PreviousExcessInputs | None = None,
    customer: str | None = None,
) -> ExcessBuildResult:
    item_data = item_data or ItemWorkbookData(metadata_by_part={}, prices_by_part={})
    layout = _layout_for_customer(customer, include_excess_input_columns)
    customer_code = customer.upper() if customer else None
    carry_inputs = previous_excess_inputs if customer_code in {"RAKEN", "AVTC"} else None
    lookup_warnings: Counter[str] = Counter()
    warning_examples: dict[str, list[str]] = defaultdict(list)
    demand_labels = _demand_labels(shortage_records)
    records_by_part = {record.part_no: record for record in shortage_records}
    known_parts = set(records_by_part)
    active_parts = {
        record.part_no
        for record in shortage_records
        if record.overshortage != 0 or record.wo_demand != 0 or record.open_po != 0
    }
    assigned_parts: set[str] = set()
    prepared_rows: list[_PreparedExcessRow] = []
    part_column_count = 1

    for record in shortage_records:
        if record.part_no in assigned_parts:
            continue

        full_group = _shortage_aware_group(
            record.part_no,
            alternates,
            records_by_part,
            known_parts,
            active_parts,
            assigned_parts,
        )

        displayed_parts = _ordered_parts(full_group)

        overshortage_values = [
            _metric(records_by_part, part, "overshortage") for part in displayed_parts
        ]
        wo_values_by_period = [
            [_demand_metric(records_by_part, part, index) for part in displayed_parts]
            for index in range(len(demand_labels))
        ]
        open_po_values = [_metric(records_by_part, part, "open_po") for part in displayed_parts]

        overshortage_total = sum(overshortage_values)
        wo_totals = [sum(values) for values in wo_values_by_period]
        open_po_total = sum(open_po_values)

        if overshortage_total == 0 and all(total == 0 for total in wo_totals) and open_po_total == 0:
            assigned_parts.update(part for part in displayed_parts if part in records_by_part)
            continue

        main_record = _first_record(displayed_parts, records_by_part)
        customer_summary = item_data.customer_summary(displayed_parts)
        model = alternates.primary_model_for(
            record.part_no,
            known_parts=known_parts,
            active_parts=active_parts,
        ) or item_data.model_summary(displayed_parts)
        model_remark = alternates.model_remark_summary(displayed_parts)
        previous_values = carry_inputs.values_for(displayed_parts) if carry_inputs else {}
        carried_values = {
            header: value for header, value in previous_values.items()
            if customer_code == "RAKEN" or header in {"reason&action", "Vendor"}
        }
        model = carried_values.get("MODEL", model)
        if customer_code == "RAKEN" and carry_inputs and carry_inputs.customer_lookup is not None:
            mapped_customer, issue = carry_inputs.customer_lookup.resolve(model)
            if issue is None:
                customer_summary = mapped_customer
            else:
                lookup_warnings[issue] += 1
                example = str(model) if model not in (None, "") else "（空白）"
                if example not in warning_examples[issue] and len(warning_examples[issue]) < 5:
                    warning_examples[issue].append(example)

        prepared_rows.append(
            _PreparedExcessRow(
                parts=displayed_parts,
                main_record=main_record,
                price=item_data.price_for(main_record.part_no),
                overshortage_values=overshortage_values,
                wo_values_by_period=wo_values_by_period,
                open_po_values=open_po_values,
                overshortage_total=overshortage_total,
                wo_totals=wo_totals,
                open_po_total=open_po_total,
                customer=customer_summary,
                model=model,
                model_remark=model_remark,
                carried_values=carried_values,
            )
        )
        part_column_count = max(part_column_count, len(displayed_parts))
        assigned_parts.update(part for part in displayed_parts if part in records_by_part)

    row1_headers = build_row1_headers(
        part_column_count,
        demand_labels,
        customer=customer,
        include_excess_input_columns=layout.include_excess_input_columns,
    )
    row2_headers = build_row2_headers(
        part_column_count,
        demand_labels,
        customer=customer,
        include_excess_input_columns=layout.include_excess_input_columns,
    )
    rows = [
        _build_output_row(
            row,
            part_column_count,
            demand_labels,
            layout,
            previous_excess_totals,
            row_number=index + 3,
            row2_headers=row2_headers,
        )
        for index, row in enumerate(prepared_rows)
    ]

    return ExcessBuildResult(
        rows=rows,
        row1_headers=row1_headers,
        row2_headers=row2_headers,
        part_column_count=part_column_count,
        input_part_count=len(shortage_records),
        output_row_count=len(rows),
        demand_labels=demand_labels,
        warnings=tuple(
            f"Customer 查表：{issue} {count} 列，保留原有 Customer；例如：{'、'.join(warning_examples[issue])}"
            for issue, count in lookup_warnings.items()
        ),
    )


def build_row1_headers(
    part_column_count: int,
    demand_labels: tuple[str, ...],
    *,
    include_excess_input_columns: bool = False,
    customer: str | None = None,
) -> list[Any]:
    layout = _layout_for_customer(customer, include_excess_input_columns)
    headers: list[Any] = [None] * (part_column_count + len(layout.info_headers))
    headers += ["Total"] + [None] * part_column_count
    for label in demand_labels:
        headers += [_group_header("Total", label, demand_labels)] + [None] * part_column_count
    headers += ["Total"] + [None] * part_column_count
    headers += ["Total"]
    headers += [None] * layout.blank_columns_before_excess
    for index, label in enumerate(demand_labels):
        headers += [None] * len(
            _excess_input_headers(
                include_excess_input_columns=layout.include_excess_input_columns,
                is_first_group=index == 0,
            )
        )
        headers += [_group_header("Excess", label, demand_labels)] + [None] * (len(layout.excess_metric_headers) - 1)
        if index == 0:
            headers += [None] * len(layout.after_first_excess_headers)
    headers += [None] * len(layout.trailing_headers)
    return headers


def build_row2_headers(
    part_column_count: int,
    demand_labels: tuple[str, ...],
    *,
    include_excess_input_columns: bool = False,
    customer: str | None = None,
) -> list[str]:
    layout = _layout_for_customer(customer, include_excess_input_columns)
    headers = (
        _part_headers(part_column_count)
        + list(layout.info_headers)
        + ["Overshortage"]
        + [f"Overshortage{index}" for index in range(1, part_column_count + 1)]
    )
    for label in demand_labels:
        headers += [_demand_total_header(label, demand_labels)]
        headers += [
            _demand_part_header(label, index, demand_labels)
            for index in range(1, part_column_count + 1)
        ]
    headers += ["Openpo"] + [f"Open po{index}" for index in range(1, part_column_count + 1)]
    headers += ["Riskbuy"]
    headers += [""] * layout.blank_columns_before_excess
    for index, label in enumerate(demand_labels):
        headers += _excess_input_headers(
            include_excess_input_columns=layout.include_excess_input_columns,
            is_first_group=index == 0,
        )
        headers += [_excess_header(header, label, demand_labels) for header in layout.excess_metric_headers]
        if index == 0:
            headers += list(layout.after_first_excess_headers)
    headers += list(layout.trailing_headers)
    return headers


def _excess_input_headers(
    *,
    include_excess_input_columns: bool,
    is_first_group: bool,
) -> list[str]:
    if not include_excess_input_columns:
        return []
    headers = [OTHER_DEMAND_HEADER] if is_first_group else []
    return headers + PO_OUTSIDE_HEADERS


def _layout_for_customer(
    customer: str | None,
    include_excess_input_columns: bool,
) -> ExcessLayout:
    customer_code = customer.upper() if customer else None
    if customer_code == "AVTC":
        return ExcessLayout(
            info_headers=tuple(AVTC_INFO_HEADERS),
            trailing_headers=tuple(AVTC_TRAILING_HEADERS),
            excess_metric_headers=tuple(AVTC_EXCESS_METRIC_HEADERS),
            include_excess_input_columns=False,
            blank_columns_before_excess=2,
        )
    return ExcessLayout(
        info_headers=tuple(INFO_HEADERS),
        trailing_headers=tuple(
            header for header in TRAILING_HEADERS
            if customer_code != "RAKEN" or header not in {"Vendor", "reason&action"}
        ),
        excess_metric_headers=tuple(EXCESS_METRIC_HEADERS),
        include_excess_input_columns=include_excess_input_columns
        or customer_code == "RAKEN",
        after_first_excess_headers=("reason&action", "Vendor") if customer_code == "RAKEN" else (),
    )


def _part_headers(part_column_count: int) -> list[str]:
    return ["PartNo1"] + [
        f"替代料{index}" for index in range(2, part_column_count + 1)
    ]


def _info_values(row: _PreparedExcessRow, layout: ExcessLayout) -> list[Any]:
    values_by_header = {
        "Desc": row.main_record.description,
        "Planner": row.main_record.planner,
        "GSD": row.main_record.buyer,
        "Customer\n客户": row.customer,
        "MODEL": row.model,
        "MODELRemark\n机种": row.model_remark,
        "LT": row.main_record.lt,
        "MOQ": row.main_record.moq,
        "Price（USD)": row.price,
    }
    return [values_by_header[header] for header in layout.info_headers]


def _trailing_values(
    row: _PreparedExcessRow,
    formulas: dict[str, Any],
    previous_total: float | None,
    layout: ExcessLayout,
) -> list[Any]:
    values_by_header = {
        "Vendor": row.carried_values.get("Vendor"),
        "reason&action": row.carried_values.get("reason&action"),
        "Customer\n客户": row.customer,
        "MODEL": row.model,
        "MODELRemark\n机种": row.model_remark,
        "Category\n分类": None,
        "Improve": formulas["improve"],
        "Previous Excess Total AMT": previous_total,
        "是否呆": None,
    }
    return [values_by_header[header] for header in layout.trailing_headers]


def _metric_formula_key(header: str) -> str:
    return {
        "Excess\nstockQTY": "stock_qty",
        "Excess\nstockAmount": "stock_amount",
        "Excess\nPOQty": "po_qty",
        "Excess\nPOAmount": "po_amount",
        "Excess\nTotalAMT": "total_amount",
    }[header]


def _build_output_row(
    row: _PreparedExcessRow,
    part_column_count: int,
    demand_labels: tuple[str, ...],
    layout: ExcessLayout,
    previous_excess_totals: PreviousExcessTotals | None,
    row_number: int,
    row2_headers: list[str],
) -> list[Any]:
    previous_total = (
        previous_excess_totals.total_for(row.parts)
        if previous_excess_totals is not None
        else None
    )
    formulas = _row_formulas(
        row_number,
        row2_headers,
        part_column_count,
        demand_labels,
        include_excess_input_columns=layout.include_excess_input_columns,
        po_outside_stock_value=row.carried_values.get(PO_OUTSIDE_STOCK_HEADER),
    )

    values: list[Any] = (
        _pad(row.parts, part_column_count)
        + _info_values(row, layout)
        + [formulas["overshortage_total"]]
        + _pad(row.overshortage_values, part_column_count, fill=0)
    )

    for formula, demand_values in zip(formulas["wo_totals"], row.wo_values_by_period):
        values += [formula] + _pad(demand_values, part_column_count, fill=0)

    values += [formulas["open_po_total"]]
    values += _pad(row.open_po_values, part_column_count, fill=0)
    values += [None]
    values += [None] * layout.blank_columns_before_excess

    for index, excess_formulas in enumerate(formulas["excess_groups"]):
        values += excess_formulas["input_values"]
        values += [
            excess_formulas[_metric_formula_key(header)]
            for header in layout.excess_metric_headers
        ]
        if index == 0:
            values += [row.carried_values.get(header) for header in layout.after_first_excess_headers]

    values += _trailing_values(row, formulas, previous_total, layout)
    return values


def _row_formulas(
    row_number: int,
    row2_headers: list[str],
    part_column_count: int,
    demand_labels: tuple[str, ...],
    *,
    include_excess_input_columns: bool = False,
    po_outside_stock_value: Any = None,
) -> dict[str, Any]:
    overshortage_total = _cell_ref(row2_headers, "Overshortage", row_number)
    overshortage_first = _cell_ref(row2_headers, "Overshortage1", row_number)
    overshortage_last = _cell_ref(
        row2_headers,
        f"Overshortage{part_column_count}",
        row_number,
    )
    open_po_total = _cell_ref(row2_headers, "Openpo", row_number)
    open_po_first = _cell_ref(row2_headers, "Open po1", row_number)
    open_po_last = _cell_ref(row2_headers, f"Open po{part_column_count}", row_number)
    price = _cell_ref(row2_headers, "Price（USD)", row_number)
    previous_total = _cell_ref(row2_headers, "Previous Excess Total AMT", row_number)

    wo_total_formulas: list[str] = []
    excess_groups: list[dict[str, Any]] = []
    first_total_amount_ref = ""
    other_demand = (
        _cell_ref(row2_headers, OTHER_DEMAND_HEADER, row_number)
        if include_excess_input_columns
        else None
    )
    first_po_outside_stock = (
        _cell_ref(row2_headers, PO_OUTSIDE_STOCK_HEADER, row_number)
        if include_excess_input_columns else None
    )
    for index, label in enumerate(demand_labels):
        wo_total = _cell_ref(row2_headers, _demand_total_header(label, demand_labels), row_number)
        wo_first = _cell_ref(row2_headers, _demand_part_header(label, 1, demand_labels), row_number)
        wo_last = _cell_ref(
            row2_headers,
            _demand_part_header(label, part_column_count, demand_labels),
            row_number,
        )
        stock_qty = _cell_ref(row2_headers, _excess_header("Excess\nstockQTY", label, demand_labels), row_number)
        stock_amount = _cell_ref(row2_headers, _excess_header("Excess\nstockAmount", label, demand_labels), row_number)
        po_qty = _cell_ref(row2_headers, _excess_header("Excess\nPOQty", label, demand_labels), row_number)
        po_amount = _cell_ref(row2_headers, _excess_header("Excess\nPOAmount", label, demand_labels), row_number)
        total_amount = _cell_ref(row2_headers, _excess_header("Excess\nTotalAMT", label, demand_labels), row_number)
        input_values: list[Any] = []
        stock_qty_formula = f"=ROUND({overshortage_total}-{wo_total},0)"
        po_outside_amount = None
        if include_excess_input_columns and other_demand is not None:
            stock_qty_formula = f"=ROUND({overshortage_total}-{wo_total}-{other_demand},0)"
            po_outside_stock = _relative_cell_ref(po_qty, row_number, -3)
            po_outside_qty = _relative_cell_ref(po_qty, row_number, -2)
            po_outside_amount = _relative_cell_ref(po_qty, row_number, -1)
            if index == 0:
                input_values.append(None)
            input_values += [
                po_outside_stock_value if index == 0 else f"={first_po_outside_stock}",
                f"=IF({po_qty}>0,{po_outside_stock},IF(({po_outside_stock}+{po_qty})>0,({po_outside_stock}+{po_qty}),0))",
                f"={po_outside_qty}*{price}",
            ]
        if not first_total_amount_ref:
            first_total_amount_ref = total_amount

        total_amount_terms = f"{po_amount}+{stock_amount}"
        total_amount_blank_checks = f'{po_amount}="",{stock_amount}=""'
        if po_outside_amount is not None:
            total_amount_terms += f"+{po_outside_amount}"
            total_amount_blank_checks += f',{po_outside_amount}=""'

        wo_total_formulas.append(f"=ROUND(SUM({wo_first}:{wo_last}),0)")
        excess_groups.append(
            {
                "input_values": input_values,
                "stock_qty": stock_qty_formula,
                "stock_amount": f'=IF({price}="","",ROUND(IF({stock_qty}<0,0,{stock_qty}*{price}),0))',
                "po_qty": f"=ROUND(IF({stock_qty}>0,{open_po_total},{stock_qty}+{open_po_total}),0)",
                "po_amount": f'=IF({price}="","",ROUND(IF({po_qty}>0,{po_qty}*{price},0),0))',
                "total_amount": f"=IF(OR({total_amount_blank_checks}),0,ROUND({total_amount_terms},0))",
            }
        )

    return {
        "overshortage_total": f"=ROUND(SUM({overshortage_first}:{overshortage_last}),0)",
        "wo_totals": wo_total_formulas,
        "open_po_total": f"=ROUND(SUM({open_po_first}:{open_po_last}),0)",
        "excess_groups": excess_groups,
        "improve": f'=IF(OR({first_total_amount_ref}<0,{previous_total}=""),"",ROUND({first_total_amount_ref}-{previous_total},0))',
    }


def _cell_ref(headers: list[str], header: str, row_number: int) -> str:
    column = headers.index(header) + 1
    return f"{get_column_letter(column)}{row_number}"


def _relative_cell_ref(cell_ref: str, row_number: int, column_offset: int) -> str:
    column_letters = cell_ref.removesuffix(str(row_number))
    column = get_column_letter(column_index_from_string(column_letters) + column_offset)
    return f"{column}{row_number}"


def _shortage_aware_group(
    part_no: str,
    alternates: BomAlternates,
    records_by_part: dict[str, ShortageRecord],
    known_parts: set[str],
    active_parts: set[str],
    assigned_parts: set[str],
) -> list[str]:
    group = list(
        alternates.ordered_group_for(
            part_no,
            known_parts=known_parts,
            active_parts=active_parts,
        )
    )
    if part_no not in group:
        group.insert(0, part_no)

    visible_parts: list[str] = []
    for part in group:
        if part != part_no and part in assigned_parts:
            continue
        if part not in visible_parts:
            visible_parts.append(part)

    if part_no not in visible_parts:
        visible_parts.insert(0, part_no)

    return visible_parts


def _ordered_parts(group: list[str]) -> list[str]:
    unique_group = []
    for part in group:
        if part not in unique_group:
            unique_group.append(part)
    return unique_group


def _first_record(
    parts: list[str],
    records_by_part: dict[str, ShortageRecord],
) -> ShortageRecord:
    for part in parts:
        record = records_by_part.get(part)
        if record is not None:
            return record
    raise ValueError("Cannot build an Excess row without at least one shortage record")


def _metric(
    records_by_part: dict[str, ShortageRecord],
    part: str,
    field_name: str,
) -> int:
    record = records_by_part.get(part)
    if record is None:
        return 0
    return _round_integer(float(getattr(record, field_name)))


def _demand_metric(
    records_by_part: dict[str, ShortageRecord],
    part: str,
    demand_index: int,
) -> int:
    record = records_by_part.get(part)
    if record is None:
        return 0
    return _round_integer(float(record.wo_demand_at(demand_index)))


def _round_integer(value: float) -> int:
    if value >= 0:
        return int(value + 0.5)
    return int(value - 0.5)


def _pad(values: list[Any], length: int, fill: Any = None) -> list[Any]:
    return values[:length] + [fill] * max(0, length - len(values))


def _demand_labels(shortage_records: list[ShortageRecord]) -> tuple[str, ...]:
    for record in shortage_records:
        if record.demand_labels:
            return record.demand_labels
    return DEFAULT_DEMAND_LABELS


def _group_header(base: str, label: str, demand_labels: tuple[str, ...]) -> str:
    if len(demand_labels) == 1:
        return base
    return f"{base}\n{label}"


def _demand_total_header(label: str, demand_labels: tuple[str, ...]) -> str:
    if len(demand_labels) == 1:
        return "WO外demand"
    return f"WO外demand\n{label}"


def _demand_part_header(label: str, index: int, demand_labels: tuple[str, ...]) -> str:
    if len(demand_labels) == 1:
        return f"WO 外demand{index}"
    return f"WO 外demand{index}\n{label}"


def _excess_header(header: str, label: str, demand_labels: tuple[str, ...]) -> str:
    if len(demand_labels) == 1:
        return header
    return f"{header}\n{label}"
