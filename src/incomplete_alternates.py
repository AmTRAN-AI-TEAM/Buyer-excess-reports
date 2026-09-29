from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from bom_alternates import IncompleteAlternateGroup


HEADERS = ["Item", "Model", "Part No", "Spec"]
DEFAULT_FONT_NAME = "Microsoft YaHei"
DEFAULT_FONT_SIZE = 10


def write_incomplete_alternates_workbook(
    groups: tuple[IncompleteAlternateGroup, ...],
    output_file: Path,
) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Incomplete Alternates"
    worksheet.append(HEADERS)

    current_row = 2
    for item_number, group in enumerate(groups, start=1):
        details = group.details or ()
        if not details:
            continue

        group_start_row = current_row
        for detail in details:
            worksheet.append(
                [
                    item_number,
                    detail.model,
                    detail.part_no,
                    detail.spec,
                ]
            )
            current_row += 1

        group_end_row = current_row - 1
        if group_end_row > group_start_row:
            worksheet.merge_cells(
                start_row=group_start_row,
                start_column=1,
                end_row=group_end_row,
                end_column=1,
            )

    _apply_layout(worksheet, max(current_row - 1, 1))
    workbook.save(output_file)


def _apply_layout(worksheet: Any, last_row: int) -> None:
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    header_font = Font(name=DEFAULT_FONT_NAME, size=DEFAULT_FONT_SIZE, bold=True)
    base_font = Font(name=DEFAULT_FONT_NAME, size=DEFAULT_FONT_SIZE)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in worksheet.iter_rows(min_row=1, max_row=last_row, max_col=len(HEADERS)):
        for cell in row:
            cell.border = border
            cell.font = base_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in range(2, last_row + 1):
        worksheet.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
        worksheet.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.cell(row, 3).alignment = Alignment(vertical="center")
        worksheet.cell(row, 4).alignment = Alignment(vertical="center", wrap_text=True)

    worksheet.column_dimensions["A"].width = 8
    worksheet.column_dimensions["B"].width = 32
    worksheet.column_dimensions["C"].width = 22
    worksheet.column_dimensions["D"].width = 72
    worksheet.row_dimensions[1].height = 22
    worksheet.freeze_panes = "A2"
