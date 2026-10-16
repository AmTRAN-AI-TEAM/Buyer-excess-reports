from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from excess_builder import PreviousExcessTotals
from report_history import CustomerLookup, PreviousExcessInputs


EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
NEXT_TIME_KEYWORD = "next_time"


@dataclass(frozen=True)
class PreviousExcessSource:
    path: Path
    totals: PreviousExcessTotals
    inputs: PreviousExcessInputs | None = None
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class _PreviousColumns:
    part_columns: tuple[int, ...]
    price: int
    overshortage_details: tuple[int, ...]
    first_demand_details: tuple[int, ...]
    open_po_details: tuple[int, ...]
    other_demand: int | None
    po_outside_stock: int | None
    input_columns: dict[str, int]


def find_previous_excess_report(input_dir: Path) -> Path | None:
    matches = sorted(
        path
        for path in input_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() in EXCEL_EXTENSIONS
        and NEXT_TIME_KEYWORD in path.name.casefold()
        and not path.name.startswith("~$")
    )
    if len(matches) > 1:
        names = ", ".join(path.name for path in matches)
        raise ValueError(f"Found multiple previous Excess reports with next_time in the name: {names}")
    return matches[0] if matches else None


def load_previous_excess_source(
    path: Path,
    *,
    customer: str | None = None,
) -> PreviousExcessSource:
    workbook = load_workbook(path, data_only=False, read_only=True)
    cached_workbook = None
    try:
        worksheet = workbook["Excess"] if "Excess" in workbook.sheetnames else workbook.active
        header_row, headers = _find_header_row(worksheet)
        columns = _detect_previous_columns(headers)
        include_excess_inputs = _include_excess_input_columns(customer, columns)
        carry_previous_inputs = include_excess_inputs or bool(customer and customer.upper() == "AVTC")
        input_columns = {
            header: column for header, column in columns.input_columns.items()
            if include_excess_inputs or header in {"reason&action", "Vendor"}
        } if carry_previous_inputs else {}
        cached_workbook = load_workbook(path, data_only=True, read_only=True)
        cached_worksheet = cached_workbook[worksheet.title]
        totals_by_part: dict[str, float] = {}
        totals_by_group: dict[tuple[str, ...], float] = {}
        values_by_part: dict[str, dict[str, Any]] = {}
        values_by_group: dict[tuple[str, ...], dict[str, Any]] = {}
        value_columns = {
            columns.price, *columns.overshortage_details,
            *columns.first_demand_details, *columns.open_po_details,
        }
        value_columns.update(input_columns.values())
        if include_excess_inputs:
            if columns.other_demand is not None:
                value_columns.add(columns.other_demand)

        for row_number, (raw_row, cached_row) in enumerate(zip(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            cached_worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        ), start=header_row + 1):
            parts = _row_parts(raw_row, columns.part_columns)
            if not parts:
                continue
            row = list(raw_row)
            # Carry formula results, never references to the previous workbook's rows.
            for column in value_columns:
                value = _cell(raw_row, column)
                if isinstance(value, str) and value.startswith("="):
                    cached_value = _cell(cached_row, column)
                    if cached_value is None:
                        raise ValueError(
                            f"Previous Excess row {row_number}, column {column} has an uncached formula; "
                            "please recalculate and save it in Excel before importing"
                        )
                    row[column - 1] = cached_value
            group = tuple(sorted(set(parts)))
            if carry_previous_inputs and (parts[0] in totals_by_part or group in totals_by_group):
                raise ValueError(f"Ambiguous previous Excess match at row {row_number}: {parts[0]}")
            total = _previous_total_amount(
                row,
                columns,
                include_excess_input_columns=include_excess_inputs,
            )
            totals_by_part.setdefault(parts[0], total)
            totals_by_group.setdefault(group, total)
            if carry_previous_inputs:
                values = {header: _cell(row, column) for header, column in input_columns.items()}
                values_by_part[parts[0]] = values
                values_by_group[group] = values

        customer_lookup = None
        warnings: list[str] = []
        if include_excess_inputs:
            if "Customer" in cached_workbook.sheetnames:
                customer_lookup = CustomerLookup.from_rows(tuple(
                    cached_workbook["Customer"].iter_rows(values_only=True)
                ))
            else:
                warnings.append("上一版沒有 Customer sheet，保留本期原有 Customer 資訊。")

        return PreviousExcessSource(
            path=path,
            totals=PreviousExcessTotals(
                totals_by_part=totals_by_part,
                totals_by_group=totals_by_group,
            ),
            inputs=PreviousExcessInputs(values_by_part, values_by_group, customer_lookup) if carry_previous_inputs else None,
            warnings=tuple(warnings),
        )
    finally:
        workbook.close()
        if cached_workbook is not None:
            cached_workbook.close()


def _find_header_row(worksheet: Any) -> tuple[int, tuple[Any, ...]]:
    for number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        if any(_base_header(value) == "PartNo1" for value in row):
            return number, row
    raise ValueError("Cannot find previous Excess header row containing PartNo1")


def _detect_previous_columns(headers: tuple[Any, ...]) -> _PreviousColumns:
    bases = [_base_header(header) for header in headers]

    part_columns = tuple(
        index
        for index, base in enumerate(bases, start=1)
        if base.startswith("PartNo") or base.startswith("替代料")
    )
    first_demand_total = _require_base_column(bases, "WO外demand")
    # In the old layout PO外備料QTY was computed. Prefer the old raw input
    # PO外備料 when present, otherwise use the newly named PO外備料 QTY.
    po_outside_stock = (
        _optional_base_column(bases, "PO外備料")
        or _optional_base_column(bases, "PO外備料 QTY")
    )
    input_columns = {
        header: column
        for header in ("MODEL", "reason&action", "Vendor")
        for column in [_optional_base_column(bases, header)]
        if column is not None
    }
    if po_outside_stock is not None:
        input_columns["PO外備料 QTY"] = po_outside_stock

    return _PreviousColumns(
        part_columns=part_columns,
        price=_require_base_column(bases, "Price（USD)"),
        overshortage_details=tuple(
            index
            for index, base in enumerate(bases, start=1)
            if base.startswith("Overshortage") and base != "Overshortage"
        ),
        first_demand_details=_following_detail_columns(bases, first_demand_total, "WO 外demand"),
        open_po_details=tuple(
            index
            for index, base in enumerate(bases, start=1)
            if base.startswith("Open po")
        ),
        other_demand=_optional_base_column(bases, "others demand"),
        po_outside_stock=po_outside_stock,
        input_columns=input_columns,
    )


def _previous_total_amount(
    row: tuple[Any, ...],
    columns: _PreviousColumns,
    *,
    include_excess_input_columns: bool,
) -> float:
    price = _number_or_none(_cell(row, columns.price))
    overshortage_total = _sum_columns(row, columns.overshortage_details)
    demand_total = _sum_columns(row, columns.first_demand_details)
    open_po_total = _sum_columns(row, columns.open_po_details)
    other_demand = (
        _number(_cell(row, columns.other_demand))
        if include_excess_input_columns and columns.other_demand is not None
        else 0.0
    )

    stock_qty = _round_integer(overshortage_total - demand_total - other_demand)
    if price is None:
        return 0.0

    stock_amount = _round_integer(max(stock_qty, 0) * price)
    po_qty = _round_integer(open_po_total if stock_qty > 0 else stock_qty + open_po_total)
    po_amount = _round_integer(po_qty * price if po_qty > 0 else 0)
    po_outside_stock = (
        _number(_cell(row, columns.po_outside_stock))
        if include_excess_input_columns and columns.po_outside_stock is not None
        else 0.0
    )
    po_outside_qty = (
        po_outside_stock
        if po_qty > 0
        else max(po_outside_stock + po_qty, 0)
    )
    po_outside_amount = _round_integer(po_outside_qty * price)
    return float(_round_integer(stock_amount + po_amount + po_outside_amount))


def _include_excess_input_columns(
    customer: str | None,
    columns: _PreviousColumns,
) -> bool:
    if customer:
        return customer.upper() == "RAKEN"
    return columns.other_demand is not None or columns.po_outside_stock is not None


def _row_parts(
    row: tuple[Any, ...],
    part_columns: tuple[int, ...],
) -> list[str]:
    return [
        text
        for column in part_columns
        for text in [_text_or_none(_cell(row, column))]
        if text
    ]


def _sum_columns(row: tuple[Any, ...], columns: tuple[int, ...]) -> float:
    return sum(_number(_cell(row, column)) for column in columns)


def _cell(row: tuple[Any, ...], column: int) -> Any:
    index = column - 1
    if index >= len(row):
        return None
    return row[index]


def _following_detail_columns(
    bases: list[str],
    start_column: int,
    detail_base_prefix: str,
) -> tuple[int, ...]:
    columns: list[int] = []
    for index in range(start_column + 1, len(bases) + 1):
        if not bases[index - 1].startswith(detail_base_prefix):
            break
        columns.append(index)
    return tuple(columns)


def _require_base_column(bases: list[str], base_header: str) -> int:
    column = _optional_base_column(bases, base_header)
    if column is None:
        raise ValueError(f"Cannot find required previous Excess column: {base_header}")
    return column


def _optional_base_column(bases: list[str], base_header: str) -> int | None:
    for index, base in enumerate(bases, start=1):
        if "".join(base.split()) == "".join(base_header.split()):
            return index
    return None


def _base_header(header: Any) -> str:
    if header in (None, ""):
        return ""
    parts = str(header).split("\n")
    if parts[0] == "Excess" and len(parts) >= 2:
        return "\n".join(parts[:2])
    return parts[0].strip()


def _number(value: Any) -> float:
    number = _number_or_none(value)
    return number if number is not None else 0.0


def _number_or_none(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, str) and value.startswith("="):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _text_or_none(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    return text or None


def _round_integer(value: float) -> int:
    if value >= 0:
        return int(value + 0.5)
    return int(value - 0.5)
