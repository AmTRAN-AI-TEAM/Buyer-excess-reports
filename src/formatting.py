from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_FONT_NAME = "Microsoft YaHei"
DEFAULT_FONT_SIZE = 9
NEGATIVE_RED_INTEGER_FORMAT = "#,##0;[Red]-#,##0"
NEGATIVE_RED_DECIMAL_FORMAT = "#,##0.00;[Red]-#,##0.00"
TOTAL_AMOUNT_SORT_FORMAT = "#,##0;;0"


def write_excess_workbook(
    rows: list[list[Any]],
    output_file: Path,
    row1_headers: list[Any],
    row2_headers: list[str],
    part_column_count: int,
) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Excess"

    worksheet.append(row1_headers)
    worksheet.append(row2_headers)
    for row in rows:
        worksheet.append(row)

    _apply_layout(worksheet, len(rows) + 2, row2_headers, part_column_count)
    workbook.save(output_file)


def _apply_layout(
    worksheet: Any,
    last_row: int,
    row2_headers: list[str],
    part_column_count: int,
) -> None:
    last_column = len(row2_headers)
    last_column_letter = get_column_letter(last_column)
    excess_start = _column_index(row2_headers, "Excess\nstockQTY")
    excess_end = _column_index(row2_headers, "Excess\nTotalAMT")

    worksheet.merge_cells(
        start_row=1,
        start_column=excess_start,
        end_row=1,
        end_column=excess_end,
    )
    worksheet.auto_filter.ref = f"A2:{last_column_letter}{max(last_row, 2)}"
    _apply_total_amount_sort(worksheet, row2_headers, last_row)
    worksheet.freeze_panes = "A3"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    group_fill = PatternFill("solid", fgColor="D9EAF7")
    base_font = Font(name=DEFAULT_FONT_NAME, size=DEFAULT_FONT_SIZE)
    header_font = Font(
        name=DEFAULT_FONT_NAME,
        size=DEFAULT_FONT_SIZE,
        color="FFFFFF",
        bold=True,
    )
    group_font = Font(name=DEFAULT_FONT_NAME, size=DEFAULT_FONT_SIZE, bold=True)
    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=max(last_row, 2),
        max_col=last_column,
    ):
        for cell in row:
            cell.border = border
            cell.font = base_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for cell in worksheet[1]:
        if cell.value is not None:
            cell.fill = group_fill
            cell.font = group_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in worksheet[2]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths_by_header = {
        "Desc": 42,
        "Planner": 14,
        "GSD": 14,
        "LT": 10,
        "MOQ": 10,
        "Price（USD)": 12,
        "Vendor": 18,
        "reason&action": 42,
        "Customer\n客户": 18,
        "MODEL": 22,
        "MODELRemark\n机种": 28,
        "Category\n分类": 16,
        "Improve": 14,
        "Previous Excess Total AMT": 18,
        "是否呆": 12,
    }

    for column in range(1, last_column + 1):
        letter = get_column_letter(column)
        header = row2_headers[column - 1]
        worksheet.column_dimensions[letter].width = widths_by_header.get(header, 12)

    for column in range(1, part_column_count + 1):
        letter = get_column_letter(column)
        worksheet.column_dimensions[letter].width = 16

    worksheet.row_dimensions[1].height = 18
    worksheet.row_dimensions[2].height = 34

    for column, header in enumerate(row2_headers, start=1):
        if _is_numeric_header(header):
            number_format = (
                NEGATIVE_RED_DECIMAL_FORMAT
                if header == "Price（USD)"
                else TOTAL_AMOUNT_SORT_FORMAT
                if header == "Excess\nTotalAMT"
                else NEGATIVE_RED_INTEGER_FORMAT
            )
            for row in range(3, last_row + 1):
                worksheet.cell(row, column).number_format = number_format

    for column in range(1, part_column_count + 1):
        for row in range(3, last_row + 1):
            worksheet.cell(row, column).number_format = "@"


def _column_index(headers: list[str], header: str) -> int:
    return headers.index(header) + 1


def _apply_total_amount_sort(
    worksheet: Any,
    row2_headers: list[str],
    last_row: int,
) -> None:
    if last_row < 3:
        return

    total_amount_column = _column_index(row2_headers, "Excess\nTotalAMT")
    total_amount_letter = get_column_letter(total_amount_column)
    worksheet.auto_filter.add_sort_condition(
        f"{total_amount_letter}3:{total_amount_letter}{last_row}",
        descending=True,
    )


def _is_numeric_header(header: str) -> bool:
    return (
        header in {
            "LT",
            "MOQ",
            "Price（USD)",
            "Overshortage",
            "WO外demand",
            "Openpo",
            "Riskbuy",
            "Excess\nstockQTY",
            "Excess\nstockAmount",
            "Excess\nPOQty",
            "Excess\nPOAmount",
            "Excess\nTotalAMT",
            "Improve",
            "Previous Excess Total AMT",
        }
        or header.startswith("Overshortage")
        or header.startswith("WO 外demand")
        or header.startswith("Open po")
    )
